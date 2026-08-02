"""Phase 4 — internal costing XLSX: build-up + totals from the frozen snapshot."""

from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.test_api_itinerary_quote import _setup, api  # noqa: F401 — `api` is a fixture


def _quote_orm(quote_id: str):
    from app.api.deps import current_org_id
    from app.db import get_session
    from app.main import app
    from app.models import Itinerary, Project, Quote

    session = next(app.dependency_overrides[get_session]())
    org_id = app.dependency_overrides[current_org_id]()
    quote = session.get(Quote, quote_id)
    assert quote is not None and quote.org_id == org_id
    return quote, session.get(Project, quote.project_id), session.get(Itinerary, quote.itinerary_id)


def test_costing_workbook_is_dynamic_and_reproduces_engine(api: TestClient) -> None:
    itinerary_id, _ = _setup(api)
    quote = api.post(f"/api/v1/itineraries/{itinerary_id}/quotes",
                     json={"buyer_state_code": "05", "fx_currency": "USD", "fx_rate": "95"}).json()

    from app.services.costing_xlsx import component_meta, render_costing_xlsx

    q, project, itinerary = _quote_orm(quote["id"])
    from app.db import get_session
    from app.main import app
    session = next(app.dependency_overrides[get_session]())
    meta = component_meta(session, itinerary.id)
    xlsx = render_costing_xlsx(q, project, itinerary, meta)
    assert xlsx[:2] == b"PK"  # a zip / xlsx

    wb = load_workbook(BytesIO(xlsx))
    assert wb.sheetnames == ["Inputs", "Cost build-up", "Rates applied", "Hotels & meal plans"]

    # The build-up is formula-driven: sell/pax and totals are Excel formulae, not
    # pre-computed floats (so editing an input recalculates everything in Excel).
    buildup = wb["Cost build-up"]
    formulas = [c.value for row in buildup.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("ROUND(" in f for f in formulas)      # sell/pax formula
    assert any("SUMPRODUCT(" in f for f in formulas)  # total cost formula

    # The engine's authoritative figures are still shown for cross-checking.
    numbers = {round(c.value, 4) for row in buildup.iter_rows() for c in row
               if isinstance(c.value, (int, float))}
    assert 336050.0 in numbers   # total cost (engine)
    assert 67277.0 in numbers    # profit (engine)
    assert 403327.0 in numbers   # package gross (engine)
    assert round(0.1252, 4) in numbers  # true margin as a fraction

    # Hotels & meal plans sheet carries a cost/pax formula per stay row.
    hotels = wb["Hotels & meal plans"]
    hotel_formulas = [c.value for row in hotels.iter_rows() for c in row
                      if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("/G" in f for f in hotel_formulas)  # rate * nights / divisor


def test_costing_endpoint(api: TestClient) -> None:
    itinerary_id, _ = _setup(api)
    quote = api.post(f"/api/v1/itineraries/{itinerary_id}/quotes",
                     json={"buyer_state_code": "05", "fx_rate": "95"}).json()
    res = api.get(f"/api/v1/quotes/{quote['id']}/costing.xlsx")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    assert res.content[:2] == b"PK"
