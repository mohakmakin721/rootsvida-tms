"""Bulk vendor/rate import — pure validation, template, and sheet parsing (no DB)."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from app.services import bulk_import as bi


# --- vendor row validation ---

def test_valid_vendor_defaults() -> None:
    clean, err = bi.validate_vendor_row({"kind": "Stay", "display_name": "Svaasa"})
    assert err is None and clean is not None
    assert clean.kind == "stay"
    assert clean.legal_name == "Svaasa"  # defaults to display_name
    assert clean.status == "prospect"    # default


def test_vendor_bad_kind_and_missing_name() -> None:
    _, err = bi.validate_vendor_row({"kind": "hotelz", "display_name": "X"})
    assert err and "kind" in err
    _, err2 = bi.validate_vendor_row({"kind": "stay", "display_name": ""})
    assert err2 and "display_name" in err2


def test_vendor_blank_row_ignored() -> None:
    clean, err = bi.validate_vendor_row({"kind": "", "display_name": ""})
    assert clean is None and err == ""


# --- rate row validation ---

def test_valid_rate_defaults() -> None:
    clean, err = bi.validate_rate_row({
        "vendor_display_name": "Svaasa", "amount": "5000",
        "valid_from": "2026-01-01", "valid_to": "2026-12-31",
    })
    assert err is None and clean is not None
    assert clean.amount == Decimal("5000")
    assert clean.meal_plan == "EP" and clean.occupancy == "single"
    assert clean.currency == "INR" and clean.rate_source == "b2b"
    assert clean.valid_from == date(2026, 1, 1)


def test_rate_bad_amount_and_dates() -> None:
    _, e1 = bi.validate_rate_row({"vendor_display_name": "X", "amount": "abc",
                                  "valid_from": "2026-01-01", "valid_to": "2026-01-02"})
    assert e1 and "number" in e1
    _, e2 = bi.validate_rate_row({"vendor_display_name": "X", "amount": "10",
                                  "valid_from": "2026-05-01", "valid_to": "2026-01-01"})
    assert e2 and "on/after" in e2
    _, e3 = bi.validate_rate_row({"vendor_display_name": "X", "amount": "-1",
                                  "valid_from": "2026-01-01", "valid_to": "2026-01-02"})
    assert e3 and ">= 0" in e3


def test_rate_accepts_datetime_and_alt_date_format() -> None:
    clean, err = bi.validate_rate_row({
        "vendor_display_name": "X", "amount": "1", "meal_plan": "map",
        "occupancy": "DOUBLE", "valid_from": "01-06-2026", "valid_to": "30-06-2026",
    })
    assert err is None and clean is not None
    assert clean.meal_plan == "MAP" and clean.occupancy == "double"
    assert clean.valid_from == date(2026, 6, 1)


# --- template + sheet parsing ---

def test_template_has_the_two_sheets_with_headers() -> None:
    wb = load_workbook(io.BytesIO(bi.build_template()))
    assert bi.VENDOR_SHEET in wb.sheetnames
    assert bi.RATE_SHEET in wb.sheetnames
    header = [c.value for c in wb[bi.VENDOR_SHEET][1]]
    assert header == bi.VENDOR_COLUMNS


def test_read_sheet_maps_headers_and_skips_blank_rows() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = bi.VENDOR_SHEET
    ws.append(bi.VENDOR_COLUMNS)
    ws.append(["stay", "Svaasa", "", "Amritsar"] + [""] * 6)
    ws.append([""] * len(bi.VENDOR_COLUMNS))  # blank → skipped
    rows = bi._read_sheet(wb, bi.VENDOR_SHEET, bi.VENDOR_COLUMNS)
    assert len(rows) == 1
    excel_row, record = rows[0]
    assert excel_row == 2  # header is row 1
    assert record["kind"] == "stay" and record["city"] == "Amritsar"
