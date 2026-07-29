"""Read a workbook sheet into verbatim staged rows.

This is the "capture, don't interpret" layer (plan §30). It turns a sheet into a
sequence of `{column_name: raw_value}` dicts, preserving each source row's real
1-based Excel row number so any staged row traces back to the exact cells it came
from. It does NOT normalise, parse prices, or map to canonical fields — that is
deliberately downstream (Milestone 6+). The only transformation is making cell
values JSON-safe, because they land in a JSONB column.
"""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ingestion.mappings import SheetSpec


@dataclass(frozen=True)
class StagedRow:
    """One source row: its 1-based Excel row number and verbatim values."""

    row_number: int
    values: dict[str, Any]


def _json_safe(value: Any) -> Any:
    """Coerce a cell value to something JSONB can store, without interpreting it.

    Dates/times become ISO strings; str/int/float/bool/None pass through
    unchanged (a phone number stored by Excel as a float stays a float — verbatim
    means faithful to the source, not tidied). Anything else is stringified.
    """
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _header_names(raw_header: tuple[Any, ...]) -> list[str]:
    """Turn a header row into column names.

    Blank header cells get a positional name (`column_3`); duplicates are
    de-duplicated with a numeric suffix so the row dict never loses a column.
    """
    names: list[str] = []
    seen: dict[str, int] = {}
    for idx, cell in enumerate(raw_header, start=1):
        base = str(cell).strip() if cell is not None and str(cell).strip() else f"column_{idx}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        names.append(base if count == 1 else f"{base}_{count}")
    return names


def read_sheet_rows(path: str | Path, spec: SheetSpec) -> Iterator[StagedRow]:
    """Yield the data rows of `spec.sheet_name` in the workbook at `path`.

    Fully-empty rows are skipped (they carry no provenance and would only be
    noise in staging). Raises `KeyError` if the sheet is absent — a missing sheet
    is a configuration error worth surfacing loudly, not a silent no-op.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if spec.sheet_name not in wb.sheetnames:
            raise KeyError(f"sheet {spec.sheet_name!r} not found in {Path(path).name}")
        ws = wb[spec.sheet_name]

        header: list[str] | None = None
        for excel_row, raw in enumerate(ws.iter_rows(values_only=True), start=1):
            if excel_row < spec.header_row:
                continue
            if excel_row == spec.header_row:
                header = _header_names(raw)
                continue
            if excel_row < spec.data_start:
                continue
            if all(cell is None for cell in raw):
                continue
            assert header is not None  # header_row < data_start guarantees this
            values = {
                header[i] if i < len(header) else f"column_{i + 1}": _json_safe(cell)
                for i, cell in enumerate(raw)
            }
            yield StagedRow(row_number=excel_row, values=values)
    finally:
        wb.close()
